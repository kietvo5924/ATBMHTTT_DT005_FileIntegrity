# -*- coding: utf-8 -*-
"""
Script Xây Dựng Slide Thuyết Trình Báo Cáo Giữa Kỳ Chuẩn Khoa Học:
Đề tài DT005 - Thẩm định Học thuật và Minh họa Thực nghiệm Bài báo Harry Halpin (2025):
"The adversary: the philosophy of cryptography" (Journal of Cybersecurity, Oxford Academic).

Quy chuẩn thiết kế:
- Khổ 16:9 widescreen (13.333 x 7.5 inch).
- Light Academic Theme (nền sáng, card slate/blue/emerald/amber, chữ sắc nét).
- Đúng 15 slide hoàn chỉnh:
  + Slide 1: Bìa chuẩn nhận diện (Thanh dọc xanh, Trường, Khoa, Báo cáo đồ án, Tên đề tài, GVHD, SV Võ Anh Kiệt).
  + Slide 2-11: Nội dung học thuật bài báo Halpin, Cơ sở triết học, Sơ đồ Timeline, Sơ đồ Shannon, Mô hình đối thủ, Snowden, 4 Câu hỏi mở, Giới hạn tái lập, Ánh xạ Paper-to-Code.
  + Slide 12: Slide chuyển tiếp tối giản: DEMO.
  + Slide 13: Kiến trúc Diễn tập Đối kháng Red Team vs Blue Team (Sơ đồ vector trực quan).
  + Slide 14: Tổng kết Báo cáo & Lộ trình Nghiên cứu Cuối kỳ.
  + Slide 15: Slide Cảm Ơn & Q&A (Đồng bộ nhận diện với Slide 1).
"""
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Thiết lập bảng màu chuẩn Light Academic
COLOR_BG_WHITE = RGBColor(255, 255, 255)
COLOR_BG_PANEL = RGBColor(248, 250, 252)       # Slate 50
COLOR_BG_CARD = RGBColor(255, 255, 255)
COLOR_BORDER = RGBColor(226, 232, 240)         # Slate 200
COLOR_BORDER_BLUE = RGBColor(191, 219, 254)    # Blue 200

COLOR_TEXT_TITLE = RGBColor(15, 23, 42)        # Slate 900
COLOR_TEXT_PRIMARY = RGBColor(30, 41, 59)      # Slate 800
COLOR_TEXT_SECONDARY = RGBColor(71, 85, 105)   # Slate 600
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)     # Slate 500

COLOR_BRAND_NAVY = RGBColor(15, 23, 42)
COLOR_BRAND_BLUE = RGBColor(30, 64, 175)       # Royal Blue
COLOR_ACCENT_BLUE = RGBColor(37, 99, 235)      # Vibrant Blue
COLOR_ACCENT_EMERALD = RGBColor(5, 150, 105)   # Emerald 600
COLOR_ACCENT_AMBER = RGBColor(217, 119, 6)     # Amber 600
COLOR_ACCENT_ROSE = RGBColor(220, 38, 38)      # Rose 600

COLOR_BADGE_BG = RGBColor(239, 246, 255)       # Blue 50
COLOR_BADGE_TEXT = RGBColor(29, 78, 216)       # Blue 700

COLOR_EMERALD_BG = RGBColor(236, 253, 245)     # Emerald 50
COLOR_EMERALD_BORDER = RGBColor(167, 243, 208) # Emerald 200
COLOR_AMBER_BG = RGBColor(255, 251, 235)       # Amber 50
COLOR_AMBER_BORDER = RGBColor(253, 230, 138)   # Amber 200


def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG_WHITE
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_badge, subtitle_text=None):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(4.5), Inches(0.35))
        badge.fill.solid()
        badge.fill.fore_color.rgb = COLOR_BADGE_BG
        badge.line.color.rgb = COLOR_BORDER_BLUE
        badge.line.width = Pt(1)
        tf_b = badge.text_frame
        tf_b.word_wrap = True
        p_b = tf_b.paragraphs[0]
        p_b.text = category_badge.upper()
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_BADGE_TEXT
        p_b.alignment = PP_ALIGN.CENTER

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.82), Inches(11.7), Inches(0.55))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(21)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_TITLE

        if subtitle_text:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.38), Inches(11.7), Inches(0.35))
            tf_s = sub_box.text_frame
            tf_s.word_wrap = True
            tf_s.margin_left = tf_s.margin_right = tf_s.margin_top = tf_s.margin_bottom = 0
            p_s = tf_s.paragraphs[0]
            p_s.text = subtitle_text
            p_s.font.size = Pt(12)
            p_s.font.color.rgb = COLOR_TEXT_SECONDARY

    def add_card(slide, left, top, width, height, bg_color=COLOR_BG_PANEL, border_color=COLOR_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        return card

    def set_shape_text(shape, lines, font_size=Pt(9.5), bold_first=True, color=COLOR_TEXT_PRIMARY):
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.04)
        for i, line in enumerate(lines):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = line
            p.font.size = font_size
            p.font.bold = bold_first if i == 0 else False
            p.font.color.rgb = color
            p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 1: BÌA BÁO CÁO GIỮA KỲ (THIẾT KẾ ĐỒNG BỘ 100% THEO ẢNH MẪU CỦA USER)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    # Dải màu xanh dương dọc bên lề trái (nhận diện thương hiệu như ảnh đính kèm)
    left_bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.8), Inches(0.18), Inches(5.9))
    left_bar1.fill.solid()
    left_bar1.fill.fore_color.rgb = COLOR_ACCENT_BLUE
    left_bar1.line.fill.background()

    # Khối tiêu đề Trường & Khoa ở phần trên
    t_inst = s1.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11.0), Inches(0.9))
    tf_inst = t_inst.text_frame
    tf_inst.word_wrap = True
    p_uni = tf_inst.paragraphs[0]
    p_uni.alignment = PP_ALIGN.CENTER
    p_uni.text = "Trường Đại học Tài nguyên và Môi trường TP. Hồ Chí Minh"
    p_uni.font.size = Pt(15)
    p_uni.font.bold = True
    p_uni.font.color.rgb = COLOR_TEXT_PRIMARY

    p_fac = tf_inst.add_paragraph()
    p_fac.alignment = PP_ALIGN.CENTER
    p_fac.space_before = Pt(4)
    p_fac.text = "Khoa Công nghệ thông tin"
    p_fac.font.size = Pt(14)
    p_fac.font.bold = True
    p_fac.font.color.rgb = COLOR_TEXT_SECONDARY

    # Tiêu đề báo cáo lớn màu xanh dương ở giữa
    t_main = s1.shapes.add_textbox(Inches(1.2), Inches(2.5), Inches(11.0), Inches(2.2))
    tf_main = t_main.text_frame
    tf_main.word_wrap = True

    p_report = tf_main.paragraphs[0]
    p_report.alignment = PP_ALIGN.CENTER
    p_report.text = "Báo cáo thuyết trình đồ án giữa kỳ"
    p_report.font.size = Pt(28)
    p_report.font.bold = True
    p_report.font.color.rgb = COLOR_ACCENT_BLUE

    p_topic = tf_main.add_paragraph()
    p_topic.alignment = PP_ALIGN.CENTER
    p_topic.space_before = Pt(12)
    p_topic.text = "Đề tài DT005: Hệ thống kiểm tra toàn vẹn tệp bằng SHA-256, HMAC và cây Merkle phân cấp động"
    p_topic.font.size = Pt(16)
    p_topic.font.bold = True
    p_topic.font.color.rgb = COLOR_TEXT_TITLE

    p_sub = tf_main.add_paragraph()
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.space_before = Pt(8)
    p_sub.text = "Khảo sát và Thẩm định học thuật bài báo: \"The adversary: the philosophy of cryptography\"\nTS. Harry Halpin — Oxford Academic, Journal of Cybersecurity, Vol. 11, No. 1, 2025 (DOI: 10.1093/cybsec/tyaf006)"
    p_sub.font.size = Pt(11)
    p_sub.font.italic = True
    p_sub.font.color.rgb = COLOR_BRAND_BLUE

    # Khối thông tin học phần, giảng viên và sinh viên ở góc dưới bên trái
    t_info = s1.shapes.add_textbox(Inches(1.4), Inches(5.1), Inches(10.0), Inches(1.6))
    tf_info = t_info.text_frame
    tf_info.word_wrap = True

    p_c = tf_info.paragraphs[0]
    p_c.text = "Học phần: An toàn và bảo mật hệ thống thông tin"
    p_c.font.size = Pt(13)
    p_c.font.color.rgb = COLOR_TEXT_PRIMARY

    p_lec = tf_info.add_paragraph()
    p_lec.space_before = Pt(4)
    p_lec.text = "Giảng viên hướng dẫn: ThS. Phạm Trọng Huynh"
    p_lec.font.size = Pt(13)
    p_lec.font.color.rgb = COLOR_TEXT_PRIMARY

    p_stu = tf_info.add_paragraph()
    p_stu.space_before = Pt(4)
    r_s1 = p_stu.add_run()
    r_s1.text = "Sinh viên thực hiện: "
    r_s1.font.size = Pt(13)
    r_s1.font.color.rgb = COLOR_TEXT_PRIMARY
    r_s2 = p_stu.add_run()
    r_s2.text = "Võ Anh Kiệt"
    r_s2.font.bold = True
    r_s2.font.size = Pt(13)
    r_s2.font.color.rgb = COLOR_BRAND_BLUE
    r_s3 = p_stu.add_run()
    r_s3.text = " — MSSV: 1150080143 — Lớp: 11_CNPM2"
    r_s3.font.size = Pt(13)
    r_s3.font.color.rgb = COLOR_TEXT_PRIMARY

    # =========================================================================
    # SLIDE 2: BỐI CẢNH, MỤC TIÊU & ĐÓNG GÓP BÀI BÁO HALPIN (SEC 1, TR. 1–2)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "BỐI CẢNH, MỤC TIÊU & BA ĐÓNG GÓP CỦA BÀI BÁO", "SECTION 1 — GIỚI THIỆU & MỤC TIÊU BÀI BÁO (TRANG 1–2)",
               "Khảo cứu nhận thức luận về Kẻ thù trong Mật mã học và Chương trình nghiên cứu tương lai")

    add_card(s2, Inches(0.8), Inches(1.85), Inches(5.7), Inches(4.8), COLOR_BG_PANEL, COLOR_BORDER)
    t = s2.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.5))
    tf = t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. Bối Cảnh Mật Mã Học & Nghịch Lý Nhận Thức Luận"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_BRAND_BLUE

    points_s2_left = [
        ("Mật mã học trong kỷ nguyên số:", "Khoa học về truyền thông bí mật khi có sự hiện diện của đối thủ có chủ đích. Internet biến không gian mạng thành nơi tương tác với vô vàn đối thủ tiềm tàng, xa lạ."),
        ("Nghịch lý nhận thức luận:", "Lý thuyết thông tin của Shannon (1948) được triết học khảo cứu sâu sắc, nhưng chính mật mã học — cái nôi khai sinh lý thuyết thông tin — lại bị triết học bỏ quên suốt nhiều thập kỷ."),
        ("Cái giá của sự trừu tượng hóa:", "Khi mật mã học biến thành các bài toán toán học thuần túy trên giấy, nó tạo ra điểm mù nguy hiểm trước các đòn tấn công vật lý trong thế giới thực.")
    ]
    for h, b in points_s2_left:
        p_it = tf.add_paragraph()
        p_it.space_before = Pt(8)
        r1 = p_it.add_run()
        r1.text = f"• {h} "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_TEXT_PRIMARY
        r2 = p_it.add_run()
        r2.text = b
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = COLOR_TEXT_SECONDARY

    add_card(s2, Inches(6.8), Inches(1.85), Inches(5.7), Inches(4.8), COLOR_EMERALD_BG, COLOR_EMERALD_BORDER)
    t = s2.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.5))
    tf = t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. Ba Đóng Góp Học Thuật Chính Của Harry Halpin"
    p.font.size = Pt(13.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_EMERALD

    points_s2_right = [
        ("Khảo sát Lịch sử & Triết học:", "Hệ thống hóa toàn diện quá trình tiến hóa mô hình đối thủ: từ Shannon (1949), PPT Adversary (1984), Dolev-Yao (1983) đến Random Oracle Model (1993)."),
        ("Phê phán Giới hạn Hình thức hóa:", "Chỉ rõ sự sụp đổ của các mô hình toán học trước thực tiễn: rò rỉ kênh kề (Kocher 1996), độ dài bản mã, rò rỉ siêu dữ liệu và giám sát hàng loạt của NSA (Snowden 2013)."),
        ("Đề xuất Chương trình Mới:", "Thiết lập 4 câu hỏi nghiên cứu mở (Section 7) nhằm mở rộng mật mã học từ một nhánh toán học biệt lập thành một khoa học an ninh hoàn chỉnh kết hợp thực nghiệm.")
    ]
    for h, b in points_s2_right:
        p_it = tf.add_paragraph()
        p_it.space_before = Pt(8)
        r1 = p_it.add_run()
        r1.text = f"• {h} "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_TEXT_PRIMARY
        r2 = p_it.add_run()
        r2.text = b
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = COLOR_TEXT_SECONDARY

    f_box = s2.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Trích dẫn: Harry Halpin (2025), Section 1: Introduction, Journal of Cybersecurity, tr. 1–2."
    p_f.font.size = Pt(10)
    p_f.font.italic = True
    p_f.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 3: CƠ SỞ TRIẾT HỌC: DRETSKE, SMITH & TÁC NHÂN (SEC 2 & 3, TR. 2–5)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "CƠ SỞ TRIẾT HỌC: THÔNG TIN, BÍ MẬT VÀ TÁC NHÂN", "SECTION 2 & 3 — CƠ SỞ TRIẾT HỌC VỀ THÔNG TIN (TRANG 2–5)",
               "Khung nhận thức luận của Fred Dretske và Brian Cantwell Smith ứng dụng trong Mật mã học")

    col_w = Inches(3.7)
    gap = Inches(0.3)
    cards_data = [
        ("1. Bản Thể Học Thông Tin", "BRIAN CANTWELL SMITH (1996)",
         COLOR_BG_PANEL, COLOR_BORDER_BLUE, COLOR_BRAND_BLUE,
         [("Dòng chảy nguyên thủy:", "Thế giới khởi đầu từ dòng chảy siêu hình không biên giới (primordial metaphysical flux)."),
          ("Phân ranh giới:", "Con người và hệ thống định hình thực tại bằng cách đăng ký đối tượng (registration of objects)."),
          ("Tác nhân xa (Distal):", "Thông tin mật mã ra đời để truyền đạt trạng thái giữa các thực thể ở xa nhau mà không thể tương tác trực tiếp.")]),
        ("2. Thuyết Thông Tin Khách Quan", "FRED DRETSKE (1981)",
         COLOR_AMBER_BG, COLOR_AMBER_BORDER, COLOR_ACCENT_AMBER,
         [("Thông tin khách quan:", "Thông tin là tính quy luật của trạng thái sự vật (state of affairs), hoàn toàn độc lập với việc có ai đọc hay không."),
          ("Thông tin và Bí mật:", "Bản rõ chứa đựng thông tin về nguồn; bản mã khi được mã hóa hoàn hảo sẽ không làm rò rỉ bất kỳ thông tin nào về nguồn."),
          ("Dòng chảy thông tin:", "Nhiệm vụ của mật mã là kiểm soát dòng chảy thông tin chỉ đi đến người nhận được ủy quyền.")]),
        ("3. Tác Nhân Có Chủ Đích", "BẢN CHẤT CỦA ĐỐI THỦ",
         COLOR_EMERALD_BG, COLOR_EMERALD_BORDER, COLOR_ACCENT_EMERALD,
         [("Tính chủ đích (Intentionality):", "Đối thủ không phải là 'nhiễu ngẫu nhiên' mà là thực thể có niềm tin, mục tiêu và ý chí khai thác lỗ hổng."),
          ("Tính tương đối của bí mật:", "Bí mật không vĩnh cửu. Không có hệ thống nào tuyệt đối; bí mật phụ thuộc vào năng lực tính toán và thời gian của đối thủ."),
          ("Public vs Common Knowledge:", "Bản mã là tri thức công cộng (public); nhưng ý nghĩa bản rõ chỉ có thể thành tri thức phổ thông nếu có khóa.")])
    ]

    for idx, (c_title, c_badge, bg_c, brd_c, text_c, items) in enumerate(cards_data):
        left_pos = Inches(0.8) + idx * (col_w + gap)
        add_card(s3, left_pos, Inches(1.85), col_w, Inches(4.8), bg_c, brd_c)
        t = s3.shapes.add_textbox(left_pos + Inches(0.2), Inches(2.0), col_w - Inches(0.4), Inches(4.5))
        tf = t.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = c_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = text_c

        sub_p = tf.add_paragraph()
        sub_p.text = c_badge
        sub_p.font.size = Pt(9.5)
        sub_p.font.bold = True
        sub_p.font.color.rgb = COLOR_TEXT_MUTED
        sub_p.space_after = Pt(8)

        for h, b in items:
            p_it = tf.add_paragraph()
            p_it.space_before = Pt(6)
            r1 = p_it.add_run()
            r1.text = f"• {h} "
            r1.font.bold = True
            r1.font.size = Pt(11)
            r1.font.color.rgb = COLOR_TEXT_PRIMARY
            r2 = p_it.add_run()
            r2.text = b
            r2.font.size = Pt(10)
            r2.font.color.rgb = COLOR_TEXT_SECONDARY

    f_box = s3.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Trích dẫn khoa học: Dretske (1981, Ref [1]); Cantwell Smith (1996, Ref [5]); Halpin (2025), tr. 2–5."
    p_f.font.size = Pt(10)
    p_f.font.italic = True
    p_f.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 4: TIẾN TRÌNH LỊCH SỬ MÔ HÌNH HÓA ĐỐI THỦ (SƠ ĐỒ 1: TIMELINE)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "TIẾN TRÌNH LỊCH SỬ MÔ HÌNH HÓA ĐỐI THỦ TRONG MẬT MÃ", "SƠ ĐỒ 1: DÒNG THỜI GIAN TIẾN HÓA MÔ HÌNH ĐỐI THỦ",
               "Sự dịch chuyển từ Lý thuyết thông tin sang Giả định độ phức tạp tính toán và Phê phán hậu Snowden")

    timeline_phases = [
        ("Giai đoạn 1 (1948–1949)", "MÔ HÌNH SHANNON",
         "Claude Shannon hình thức hóa mật mã học bằng xác suất; định nghĩa Bí mật hoàn hảo (OTP). Giả định sai lệch do 'Nhiễu ngẫu nhiên'.",
         COLOR_BG_PANEL, COLOR_BORDER_BLUE, COLOR_BRAND_BLUE),
        ("Giai đoạn 2 (1976)", "CÁCH MẠNG KHÓA CÔNG KHAI",
         "Diffie-Hellman và Merkle giải quyết bài toán phân phối khóa; chuyển gánh nặng bảo mật sang độ phức tạp tính toán một chiều.",
         COLOR_AMBER_BG, COLOR_AMBER_BORDER, COLOR_ACCENT_AMBER),
        ("Giai đoạn 3 (1982–1984)", "PPT ADVERSARY & SECURITY GAMES",
         "Goldwasser-Micali định nghĩa an toàn ngữ nghĩa và đối thủ thời gian đa thức xác suất (PPT). Chứng minh an toàn bằng quy dẫn toán học.",
         COLOR_BG_PANEL, COLOR_BORDER, COLOR_TEXT_TITLE),
        ("Giai đoạn 4 (1981–1993)", "DOLEV-YAO & RANDOM ORACLE",
         "Dolev-Yao mô hình hóa đối thủ mạng kiểm soát hoàn toàn kênh truyền. Bellare-Rogaway đề xuất Random Oracle Model (ROM) ứng dụng thực hành.",
         COLOR_EMERALD_BG, COLOR_EMERALD_BORDER, COLOR_ACCENT_EMERALD),
        ("Giai đoạn 5 (2013–2025)", "PHÊ PHÁN HẬU SNOWDEN",
         "Tiết lộ của Snowden phơi bày năng lực thực tế của NSA: Kênh kề, cửa sau chuẩn mật mã, tấn công chuỗi cung ứng và giám sát siêu dữ liệu.",
         RGBColor(254, 242, 242), RGBColor(254, 202, 202), COLOR_ACCENT_ROSE)
    ]

    for idx, (t_time, t_badge, t_desc, bg_c, brd_c, text_c) in enumerate(timeline_phases):
        top_pos = Inches(1.85) + idx * Inches(0.95)
        add_card(s4, Inches(0.8), top_pos, Inches(11.733), Inches(0.85), bg_c, brd_c)
        t = s4.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.05), Inches(11.333), Inches(0.75))
        tf = t.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r_time = p.add_run()
        r_time.text = f"{t_time} — "
        r_time.font.bold = True
        r_time.font.size = Pt(11.5)
        r_time.font.color.rgb = text_c

        r_badge = p.add_run()
        r_badge.text = f"[{t_badge}]: "
        r_badge.font.bold = True
        r_badge.font.size = Pt(11)
        r_badge.font.color.rgb = COLOR_TEXT_PRIMARY

        r_desc = p.add_run()
        r_desc.text = t_desc
        r_desc.font.size = Pt(10.5)
        r_desc.font.color.rgb = COLOR_TEXT_SECONDARY

    f_box = s4.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Sơ đồ tổng hợp dựa trên Section 3–6, Harry Halpin (2025), Journal of Cybersecurity, tr. 3–12."
    p_f.font.size = Pt(10)
    p_f.font.italic = True
    p_f.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 5: SHANNON & MẠNG PHỨC HỢP (SƠ ĐỒ 2: MÔ HÌNH TRUYỀN THÔNG)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "MÔ HÌNH TRUYỀN THÔNG CỦA SHANNON VÀ MẠNG THỰC TẾ", "SƠ ĐỒ 2: ĐỐI SÁNH MÔ HÌNH TRUYỀN THÔNG ĐIỂM-ĐIỂM VÀ MẠNG XÃ HỘI-KỸ THUẬT",
               "Khoảng cách giữa mô hình điểm-điểm đóng kín của Shannon và mạng phức hợp không đồng nhất")

    img_s5_path = os.path.join(os.path.dirname(__file__), "docs", "Slide6.drawio.png")
    if os.path.exists(img_s5_path):
        img_w = Inches(10.3)
        img_h = Inches(4.84)
        img_l = (prs.slide_width - img_w) / 2
        img_t = Inches(1.82)
        s5.shapes.add_picture(img_s5_path, img_l, img_t, img_w, img_h)

    f_box = s5.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Sơ đồ tổng hợp dựa trên Shannon (1949) và Harry Halpin (2025), Journal of Cybersecurity, tr. 3–5."
    p_f.font.size = Pt(10)
    p_f.font.italic = True
    p_f.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 6: CÁC MÔ HÌNH ĐỐI THỦ HÌNH THỨC (SEC 4, TR. 5–8)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "BA MÔ HÌNH ĐỐI THỦ HÌNH THỨC TRONG MẬT MÃ HIỆN ĐẠI", "SECTION 4 — HÌNH THỨC HÓA ĐỐI THỦ (TRANG 5–8)",
               "Tiến trình toán học hóa kẻ thù: Độ phức tạp tính toán đa thức, Logic kênh truyền và Tiên tri ngẫu nhiên")

    col_w7 = Inches(3.7)
    gap7 = Inches(0.3)
    formal_models = [
        ("1. Mô Hình PPT Adversary", "GOLDWASSER-MICALI (1984)",
         COLOR_BG_PANEL, COLOR_BORDER_BLUE, COLOR_BRAND_BLUE,
         [("Khái niệm:", "Đối thủ là thuật toán Turing xác suất chạy trong thời gian đa thức (Probabilistic Polynomial-Time)."),
          ("Security Games:", "Định nghĩa an toàn thông qua các trò chơi thách đấu. Lợi thế Adv(A) phải là hàm không đáng kể (negligible)."),
          ("Quy dẫn toán học:", "Chứng minh an toàn bằng cách quy dẫn: bẻ gãy hệ mật tương đương giải bài toán khó (phân tích thừa số nguyên, logarit rời rạc).")]),
        ("2. Mô Hình Đối Thủ Mạng Logic", "DOLEV-YAO (1983)",
         COLOR_BG_PANEL, COLOR_BORDER, COLOR_TEXT_TITLE,
         [("Khái niệm:", "Đối thủ kiểm soát toàn bộ kênh truyền mạng: có quyền đọc, chặn, sửa đổi, xóa, tráo đổi và phát lại gói tin."),
          ("Hộp đen mật mã:", "Mô hình Dolev-Yao xem mật mã như các hàm đại số hoàn hảo (Perfect Black-box). Kẻ thù không thể giải mã nếu không có khóa."),
          ("Ứng dụng:", "Kiểm chứng hình thức các giao thức mạng phức tạp thông qua các công cụ logic tự động.")]),
        ("3. Mô Hình Tiên Tri Ngẫu Nhiên", "BELLARE-ROGAWAY (1993)",
         COLOR_AMBER_BG, COLOR_AMBER_BORDER, COLOR_ACCENT_AMBER,
         [("Khái niệm:", "Giả định hàm băm thực tế hoạt động như một hàm ngẫu nhiên lý tưởng (Random Oracle Model - ROM) trả về chuỗi ngẫu nhiên đồng đều."),
          ("Vai trò thực tế:", "Cho phép xây dựng các hệ mật khóa công khai và chữ ký số hiệu năng cao có chứng minh an toàn."),
          ("Giới hạn triết học:", "Trong thế giới thực, không có hàm băm cụ thể nào (kể cả SHA-256) có thể tương đương toán học với một Random Oracle thực sự.")])
    ]

    for idx, (m_title, m_badge, bg_c, brd_c, text_c, items) in enumerate(formal_models):
        left_pos = Inches(0.8) + idx * (col_w7 + gap7)
        add_card(s6, left_pos, Inches(1.85), col_w7, Inches(4.8), bg_c, brd_c)
        t = s6.shapes.add_textbox(left_pos + Inches(0.2), Inches(2.0), col_w7 - Inches(0.4), Inches(4.5))
        tf = t.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = m_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = text_c

        sub_p = tf.add_paragraph()
        sub_p.text = m_badge
        sub_p.font.size = Pt(9.5)
        sub_p.font.bold = True
        sub_p.font.color.rgb = COLOR_TEXT_MUTED
        sub_p.space_after = Pt(8)

        for h, b in items:
            p_it = tf.add_paragraph()
            p_it.space_before = Pt(6)
            r1 = p_it.add_run()
            r1.text = f"• {h} "
            r1.font.bold = True
            r1.font.size = Pt(11)
            r1.font.color.rgb = COLOR_TEXT_PRIMARY
            r2 = p_it.add_run()
            r2.text = b
            r2.font.size = Pt(10)
            r2.font.color.rgb = COLOR_TEXT_SECONDARY

    f_box = s6.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Trích dẫn: Goldwasser-Micali (1984); Dolev-Yao (1983); Bellare-Rogaway (1993); Halpin (2025), tr. 5–8."
    p_f.font.size = Pt(10)
    p_f.font.italic = True
    p_f.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 7: GIỚI HẠN CỦA HÌNH THỨC HÓA (SƠ ĐỒ 3: KHOẢNG CÁCH VỚI THỰC TẾ)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "GIỚI HẠN CỦA SỰ HÌNH THỨC HÓA TRONG MẬT MÃ HỌC", "SECTION 5 — GIỚI HẠN TRỪU TƯỢNG HÓA (TRANG 8–10)",
               "Khoảng cách giữa các giả định toán học lý tưởng và thế giới vật lý thực tế")

    flow_steps = [
        ("BƯỚC 1: GIẢ ĐỊNH TOÁN HỌC HÌNH THỨC",
         "Hộp đen mật mã lý tưởng, thời gian thực thi là vô hình, độ dài bản mã công khai không rò rỉ nội dung, kênh truyền điểm-điểm độc lập.",
         COLOR_BG_PANEL, COLOR_BORDER_BLUE, COLOR_BRAND_BLUE),
        ("BƯỚC 2: CÁC YẾU TỐ VẬT LÝ BỊ TRỪU TƯỢNG HÓA LƯỢC BỎ",
         "Thời gian vi xử lý CPU, dao động điện năng tiêu thụ, phát xạ điện từ, kiến trúc bộ nhớ đệm (Cache), lỗi lập trình phần mềm.",
         COLOR_AMBER_BG, COLOR_AMBER_BORDER, COLOR_ACCENT_AMBER),
        ("BƯỚC 3: CON ĐƯỜNG KHAI THÁC CỦA ĐỐI THỦ THẾ GIỚI THỰC",
         "Tấn công kênh kề (Timing Attack - Kocher 1996), phân tích lưu lượng siêu dữ liệu, tấn công tiền ảnh cấp hai cấu trúc (Node-Leaf Confusion).",
         RGBColor(254, 242, 242), RGBColor(254, 202, 202), COLOR_ACCENT_ROSE),
        ("BƯỚC 4: KẾT LUẬN NHẬN THỨC LUẬN CỦA HARRY HALPIN",
         "Cái giá của sự trừu tượng hóa là làm cho mật mã học mù quáng trước chính những năng lực thực tế của đối thủ nằm ngoài mô hình lý thuyết.",
         COLOR_EMERALD_BG, COLOR_EMERALD_BORDER, COLOR_ACCENT_EMERALD)
    ]

    for idx, (step_t, step_d, bg_c, brd_c, text_c) in enumerate(flow_steps):
        top_pos = Inches(1.85) + idx * Inches(1.18)
        add_card(s7, Inches(0.8), top_pos, Inches(11.733), Inches(1.05), bg_c, brd_c)
        t = s7.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.08), Inches(11.333), Inches(0.9))
        tf = t.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = step_t
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = text_c

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(3)
        p_d.text = step_d
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = COLOR_TEXT_PRIMARY

    f_box = s7.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Sơ đồ tổng hợp dựa trên Kocher (1996) và Halpin (2025), Section 5, tr. 8–10."
    p_f.font.size = Pt(10)
    p_f.font.italic = True
    p_f.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 8: TIẾT LỘ CỦA SNOWDEN & ĐỐI THỦ THỂ CHẾ (SEC 6, TR. 10–12)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "SNOWDEN VÀ ĐỐI THỦ THỂ CHẾ (THE REAL ADVERSARY)", "SECTION 6 — TIẾT LỘ CỦA SNOWDEN (TRANG 10–12)",
               "Khủng hoảng mô hình an ninh trước đối thủ giám sát hàng loạt quy mô quốc gia")

    col_w9 = Inches(3.7)
    gap9 = Inches(0.3)
    snowden_pillars = [
        ("1. Khủng Hoảng Thực Tiễn", "TIẾT LỘ SNOWDEN (2013)",
         COLOR_BG_PANEL, COLOR_BORDER_BLUE, COLOR_BRAND_BLUE,
         [("Sự kiện bước ngoặt:", "Các tài liệu Edward Snowden tiết lộ về NSA đã kéo mật mã học vào một cuộc khủng hoảng thực tiễn sâu sắc."),
          ("Không phá toán học:", "NSA không cần giải các bài toán toán học khó của RSA hay AES mà tấn công trực tiếp vào toàn bộ hệ sinh thái xung quanh."),
          ("Mô hình thực tế:", "Đối thủ bị động toàn cầu (Global Passive Adversary) và tấn công can thiệp vào đường truyền là mối đe dọa hoàn toàn có thật.")]),
        ("2. Năng Lực Đối Thủ Thể Chế", "VƯỢT NGOÀI HÌNH THỨC",
         COLOR_AMBER_BG, COLOR_AMBER_BORDER, COLOR_ACCENT_AMBER,
         [("Giám sát Siêu dữ liệu:", "NSA tập trung khai thác Metadata (ai liên lạc với ai, thời gian, kích thước) để dựng lại mạng xã hội mà không cần bẻ khóa bản mã."),
          ("Tấn công Chuỗi cung ứng:", "Chèn cửa sau vào phần cứng, làm suy yếu chuẩn mật mã quốc tế (tiêu chuẩn NIST Dual_EC_DRBG bị cài cửa sau ngầm)."),
          ("Áp lực Thể chế:", "Cưỡng chế pháp lý các tập đoàn công nghệ lớn bàn giao khóa bí mật và quyền truy cập dữ liệu máy chủ.")]),
        ("3. Tính Tương Đối Của Kẻ Thù", "BẢN CHẤT NHẬN THỨC LUẬN",
         COLOR_EMERALD_BG, COLOR_EMERALD_BORDER, COLOR_ACCENT_EMERALD,
         [("Khái niệm tương đối:", "Khái niệm 'Kẻ thù' mang tính vị thế: Đối với người dân, NSA là đối thủ giám sát; nhưng đối với NSA, Snowden lại là đối thủ làm rò rỉ bí mật."),
          ("Dân chủ hóa mật mã:", "Mật mã khóa công khai từng giúp cá nhân tự vệ, nhưng công cụ gián điệp của NSA nay đã lan rộng sang tội phạm mạng và các nhóm phi nhà nước."),
          ("Hệ quả khoa học:", "Mật mã học buộc phải từ bỏ giả định kênh truyền sạch để nghiên cứu phòng chống cửa sau (Cliptography).")])
    ]

    for idx, (s_title, s_badge, bg_c, brd_c, text_c, items) in enumerate(snowden_pillars):
        left_pos = Inches(0.8) + idx * (col_w9 + gap9)
        add_card(s8, left_pos, Inches(1.85), col_w9, Inches(4.8), bg_c, brd_c)
        t = s8.shapes.add_textbox(left_pos + Inches(0.2), Inches(2.0), col_w9 - Inches(0.4), Inches(4.5))
        tf = t.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = s_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = text_c

        sub_p = tf.add_paragraph()
        sub_p.text = s_badge
        sub_p.font.size = Pt(9.5)
        sub_p.font.bold = True
        sub_p.font.color.rgb = COLOR_TEXT_MUTED
        sub_p.space_after = Pt(8)

        for h, b in items:
            p_it = tf.add_paragraph()
            p_it.space_before = Pt(6)
            r1 = p_it.add_run()
            r1.text = f"• {h} "
            r1.font.bold = True
            r1.font.size = Pt(11)
            r1.font.color.rgb = COLOR_TEXT_PRIMARY
            r2 = p_it.add_run()
            r2.text = b
            r2.font.size = Pt(10)
            r2.font.color.rgb = COLOR_TEXT_SECONDARY

    f_box = s8.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Trích dẫn: Harry Halpin (2025), Section 6: Snowden and the real adversary, tr. 10–12."
    p_f.font.size = Pt(10)
    p_f.font.italic = True
    p_f.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 9: BỐN CÂU HỎI NGHIÊN CỨU MỞ CỦA HALPIN (SEC 7, TR. 12–14)
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9)
    add_header(s9, "CHƯƠNG TRÌNH NGHIÊN CỨU TƯƠNG LAI: BỐN CÂU HỎI MỞ", "SECTION 7 — KẾT LUẬN & CHƯƠNG TRÌNH NGHIÊN CỨU (TRANG 12–14)",
               "Đề xuất hòa giải giữa trừu tượng hóa toán học và dữ liệu thực nghiệm thế giới thực")

    q_boxes = [
        ("Câu hỏi 1: Mạng Xã Hội Động của Truyền Thông Bí Mật",
         "Can we study dynamic social networks of secret communication modeled on the empirical behavior of secret keepers?",
         "Nghiên cứu tổ chức thông tin dưới dạng mạng xã hội động giữa những người giữ bí mật. Kiểm tra xem các bất biến chính sách bảo vệ bí mật có được duy trì qua các biến động cấu trúc mạng và các sự cố Byzantine hay không."),
        ("Câu hỏi 2: Mô Hình Hóa Năng Lực và Liên Minh Đối Thủ Bằng Thống Kê",
         "By incorporating statistics, can new models of the adversary take into account more accurate data about their capabilities?",
         "Sử dụng thống kê thực nghiệm và mô phỏng Monte Carlo để mô hình hóa chính xác hơn năng lực đối thủ; phân tích kinh tế chính trị của các liên minh giám sát (liên minh Five Eyes, Anonymous, quan hệ nhà nước - cá nhân)."),
        ("Câu hỏi 3: Ứng Dụng Phương Pháp Liên Ngành & Hệ Thống Phức Hợp",
         "Can cryptography incorporate interdisciplinary methods and insights that can capture the complexity of secure communication?",
         "Áp dụng các công cụ từ lý thuyết hệ thống phức hợp, động lực học mạng và dịch tễ học để mô hình hóa sự lan truyền thông tin, nhận diện các điểm chuyển pha khi bị tấn công."),
        ("Câu hỏi 4: Triết Học Giúp Tái Cấu Trúc Các Giả Định Ngầm Định",
         "Will philosophy let us make explicit and then reformulate the hidden assumptions of cryptographic communication, and vice versa?",
         "Vận dụng triết học để chất vấn các giả định nền tảng của Shannon; xem xét ranh giới động giữa con người và môi trường kỹ thuật số theo Giả thuyết Tâm trí Mở rộng (Extended Mind Hypothesis).")
    ]

    for idx, (q_title, q_en, q_desc) in enumerate(q_boxes):
        top_pos = Inches(1.85) + idx * Inches(1.18)
        add_card(s9, Inches(0.8), top_pos, Inches(11.733), Inches(1.05), COLOR_BG_PANEL, COLOR_BORDER_BLUE)
        t = s9.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.06), Inches(11.333), Inches(0.92))
        tf = t.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = q_title
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_BRAND_BLUE

        p_en = tf.add_paragraph()
        p_en.text = f"\"{q_en}\""
        p_en.font.size = Pt(9.5)
        p_en.font.italic = True
        p_en.font.color.rgb = COLOR_TEXT_MUTED

        p_d = tf.add_paragraph()
        p_d.space_before = Pt(2)
        p_d.text = q_desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = COLOR_TEXT_PRIMARY

    f_box = s9.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Trích dẫn nguyên văn 4 câu hỏi nghiên cứu mở: Harry Halpin (2025), Section 7, tr. 12–14."
    p_f.font.size = Pt(10)
    p_f.font.italic = True
    p_f.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 10: PHƯƠNG PHÁP LUẬN & GIỚI HẠN TÁI LẬP BÀI BÁO
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10)
    add_header(s10, "PHƯƠNG PHÁP LUẬN & KHẢ NĂNG TÁI LẬP BÀI BÁO", "ĐÁNH GIÁ PHƯƠNG PHÁP LUẬN & HƯỚNG TIẾP CẬN THỰC NGHIỆM",
               "Làm rõ tính khả thi tái lập đối với một công trình triết học và định hướng thí nghiệm độc lập")

    add_card(s10, Inches(0.8), Inches(1.85), Inches(5.7), Inches(4.8), COLOR_AMBER_BG, COLOR_AMBER_BORDER)
    t = s10.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.5))
    tf = t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. Đánh Giá Khả Năng Tái Lập (Replication Feasibility)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_AMBER

    replicate_pts = [
        ("Bản chất công trình:", "Bài báo của TS. Harry Halpin là một luận văn tổng quan triết học và phân tích nhận thức luận, không phải là một công trình kỹ thuật phần mềm."),
        ("Không có tài nguyên phần mềm:", "Bài báo HOÀN TOÀN KHÔNG có mã nguồn triển khai, không có tập dữ liệu thực nghiệm, không có các bảng đo thời gian (ms) hay dung lượng benchmark."),
        ("Phương pháp thẩm định:", "Nhóm sinh viên thực hiện thẩm định học thuật (Scholarly Audit) đối chiếu toàn bộ 17 trang bản in PDF, 7 mục nội dung và 134 tài liệu tham khảo."),
        ("Tính trung thực học thuật:", "Không bịa đặt số liệu gốc của tác giả; không tuyên bố 'tái lập 100% benchmark bài báo' khi tác giả không công bố benchmark.")
    ]
    for h, b in replicate_pts:
        p_it = tf.add_paragraph()
        p_it.space_before = Pt(8)
        r1 = p_it.add_run()
        r1.text = f"• {h} "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_TEXT_PRIMARY
        r2 = p_it.add_run()
        r2.text = b
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = COLOR_TEXT_SECONDARY

    add_card(s10, Inches(6.8), Inches(1.85), Inches(5.7), Inches(4.8), COLOR_EMERALD_BG, COLOR_EMERALD_BORDER)
    t = s10.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.5))
    tf = t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. Định Hướng Thí Nghiệm Minh Họa Độc Lập"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_EMERALD

    independent_pts = [
        ("Mục tiêu thực nghiệm:", "Sinh viên tự xây dựng hệ thống kiểm tra toàn vẹn tệp nhằm chuyển hóa các luận điểm lý thuyết trừu tượng của Halpin thành các kịch bản quan sát được."),
        ("Minh họa trực quan:", "Thiết kế các thí nghiệm minh họa: sự bất lực của mã tuyến tính CRC32 trước kẻ thù toán học; can thiệp đường truyền Dolev-Yao; rò rỉ siêu dữ liệu kênh truyền."),
        ("Phân định ranh giới rõ ràng:", "Mã nguồn và hệ thống kiểm tra toàn vẹn là đóng góp kỹ thuật độc lập của sinh viên; không được nhầm lẫn là thuật toán do Halpin đề xuất."),
        ("Báo cáo khoa học nghiêm túc:", "Tất cả kết quả đo đạc được lưu vết kiểm toán tại file MIDTERM_CODE_AUDIT.md.")
    ]
    for h, b in independent_pts:
        p_it = tf.add_paragraph()
        p_it.space_before = Pt(8)
        r1 = p_it.add_run()
        r1.text = f"• {h} "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_TEXT_PRIMARY
        r2 = p_it.add_run()
        r2.text = b
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = COLOR_TEXT_SECONDARY

    f_box = s10.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Báo cáo phân định phương pháp luận của nhóm nghiên cứu đề tài DT005."
    p_f.font.size = Pt(10)
    p_f.font.italic = True
    p_f.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 11: MINH HỌA PAPER-TO-CODE ĐỘC LẬP & GIỚI HẠN DIỄN GIẢI
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_bg(s11)
    add_header(s11, "MINH HỌA PAPER-TO-CODE ĐỘC LẬP & GIỚI HẠN DIỄN GIẢI", "SƠ ĐỒ 4: PHÂN ĐỊNH LÝ THUYẾT & THÍ NGHIỆM ĐỘC LẬP",
               "Đối sánh minh bạch giữa các luận điểm bài báo và thí nghiệm minh họa do sinh viên xây dựng")

    col_w11 = Inches(3.7)
    gap11 = Inches(0.3)

    examples_data = [
        ("1. NHIỄU vs KẺ THÙ CHỦ ĐÍCH",
         "Mật mã học khác lý thuyết Shannon: đối thủ có chủ đích, không phải nhiễu ngẫu nhiên. Mã toàn vẹn tuyến tính CRC32 hoàn toàn bất lực.",
         "Giải hệ phương trình bù đại số GF(2), tìm 4 byte giả mạo qua mặt CRC32 100%. Hàm băm SHA-256 phát hiện thay đổi ngay lập tức.",
         "Chỉ minh họa tính phi tuyến của hàm băm; không chứng minh an toàn toàn diện cho toàn bộ hệ thống."),
        ("2. NHẬP NHẰNG NÚT LÁ MERKLE",
         "Ghép chuỗi thô H(L || R) khiến kẻ thù có thể đánh tráo nút nhánh nội bộ thành nút lá, gây va chạm cấu trúc tầng cây.",
         "Cài đặt phân tách miền chuẩn RFC 6962: gắn tiền tố 0x00 cho nút lá và 0x01 cho nút nhánh để triệt tiêu hoàn toàn nhập nhằng.",
         "Chỉ giải quyết va chạm cấu trúc tầng cây Merkle; không loại trừ va chạm ngẫu nhiên 2^256 của SHA-256."),
        ("3. KÊNH KỀ THỜI GIAN SO SÁNH",
         "Lý thuyết xem mật mã là hàm toán học lý tưởng, nhưng thực tế việc dừng sớm (early-return) làm rò rỉ dữ liệu qua thời gian thực thi.",
         "Đo đạc 5 vòng so sánh chuỗi tuần tự (lộ chênh lệch thời gian) so với hàm thời gian hằng số hmac.compare_digest triệt tiêu độ lệch.",
         "Đo đạc trên máy tính đơn có nhiễu hệ điều hành; chỉ minh họa ở mức thư viện phần mềm, không phải phần cứng.")
    ]

    for idx, (ex_title, ex_theory, ex_impl, ex_limit) in enumerate(examples_data):
        left_pos = Inches(0.8) + idx * (col_w11 + gap11)
        add_card(s11, left_pos, Inches(1.85), col_w11, Inches(4.8), COLOR_BG_PANEL, COLOR_BORDER)
        t = s11.shapes.add_textbox(left_pos + Inches(0.2), Inches(2.0), col_w11 - Inches(0.4), Inches(4.5))
        tf = t.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = ex_title
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_BRAND_NAVY
        p_t.space_after = Pt(10)

        # 1. Luận điểm bài báo Halpin
        p1_h = tf.add_paragraph()
        p1_h.text = "Luận điểm bài báo Halpin:"
        p1_h.font.size = Pt(10.5)
        p1_h.font.bold = True
        p1_h.font.color.rgb = COLOR_TEXT_PRIMARY

        p1_b = tf.add_paragraph()
        p1_b.text = ex_theory
        p1_b.font.size = Pt(9.5)
        p1_b.font.color.rgb = COLOR_TEXT_SECONDARY
        p1_b.space_after = Pt(8)

        # 2. Minh họa thực nghiệm DT005
        p2_h = tf.add_paragraph()
        p2_h.text = "Minh họa thực nghiệm DT005:"
        p2_h.font.size = Pt(10.5)
        p2_h.font.bold = True
        p2_h.font.color.rgb = COLOR_TEXT_PRIMARY

        p2_b = tf.add_paragraph()
        p2_b.text = ex_impl
        p2_b.font.size = Pt(9.5)
        p2_b.font.color.rgb = COLOR_TEXT_SECONDARY
        p2_b.space_after = Pt(8)

        # 3. Giới hạn diễn giải
        p3_h = tf.add_paragraph()
        p3_h.text = "Giới hạn diễn giải:"
        p3_h.font.size = Pt(10.5)
        p3_h.font.bold = True
        p3_h.font.color.rgb = COLOR_TEXT_PRIMARY

        p3_b = tf.add_paragraph()
        p3_b.text = ex_limit
        p3_b.font.size = Pt(9.5)
        p3_b.font.color.rgb = COLOR_TEXT_MUTED

    f_box = s11.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Bảng đối chiếu minh bạch do sinh viên tự xây dựng; chi tiết mã nguồn lưu tại MIDTERM_CODE_AUDIT.md."
    p_f.font.size = Pt(10)
    p_f.font.italic = True
    p_f.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 12: SLIDE CHUYỂN TIẾP TỐI GIẢN: DEMO (NỀN SÁNG TRANG NHÃ)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_bg(s12)

    add_card(s12, Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5), COLOR_BG_PANEL, COLOR_BORDER_BLUE)

    t_demo = s12.shapes.add_textbox(Inches(1.8), Inches(1.8), Inches(9.733), Inches(3.9))
    tf_demo = t_demo.text_frame
    tf_demo.word_wrap = True

    p_badge = tf_demo.paragraphs[0]
    p_badge.alignment = PP_ALIGN.CENTER
    p_badge.text = "DIỄN TẬP AN NINH MẠNG TRỰC TIẾP"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_BRAND_BLUE

    p_main = tf_demo.add_paragraph()
    p_main.alignment = PP_ALIGN.CENTER
    p_main.space_before = Pt(14)
    p_main.text = "DEMO"
    p_main.font.size = Pt(64)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_BRAND_NAVY

    p_sub = tf_demo.add_paragraph()
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.space_before = Pt(14)
    p_sub.text = "HỆ THỐNG GIÁM SÁT TOÀN VẸN TỆP & MÔ HÌNH ĐỐI KHÁNG RED TEAM VS BLUE TEAM"
    p_sub.font.size = Pt(15)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_TEXT_PRIMARY

    p_note = tf_demo.add_paragraph()
    p_note.alignment = PP_ALIGN.CENTER
    p_note.space_before = Pt(10)
    p_note.text = "Minh họa thực nghiệm các luận điểm lý thuyết: Can thiệp dữ liệu, Nhập nhằng nút lá RFC 6962, và So sánh chuỗi thời gian hằng số"
    p_note.font.size = Pt(11)
    p_note.font.italic = True
    p_note.font.color.rgb = COLOR_TEXT_SECONDARY

    p_port = tf_demo.add_paragraph()
    p_port.alignment = PP_ALIGN.CENTER
    p_port.space_before = Pt(12)
    p_port.text = "Blue Team Dashboard: http://localhost:8080   |   Red Team Adversary Engine: http://localhost:8081"
    p_port.font.size = Pt(10.5)
    p_port.font.bold = True
    p_port.font.color.rgb = COLOR_BRAND_BLUE

    # =========================================================================
    # SLIDE 13: KIẾN TRÚC DIỄN TẬP ĐỐI KHÁNG RED TEAM VS BLUE TEAM (SƠ ĐỒ HÌNH VẼ)
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_bg(s13)
    add_header(s13, "KIẾN TRÚC DIỄN TẬP ĐỐI KHÁNG: RED TEAM VS BLUE TEAM", "SƠ ĐỒ 5: KIẾN TRÚC ĐỐI KHÁNG THỰC TẾ & BẰNG CHỨNG KIỂM ĐỊNH",
               "Mô hình đối kháng thực tế 2 màn hình trên tập dữ liệu kaggle_dataset (2.483 tệp tin)")

    img_s13_path = os.path.join(os.path.dirname(__file__), "docs", "Slide14.drawio.png")
    if os.path.exists(img_s13_path):
        img_w = Inches(10.3)
        img_h = Inches(4.84)
        img_l = (prs.slide_width - img_w) / 2
        img_t = Inches(1.82)
        s13.shapes.add_picture(img_s13_path, img_l, img_t, img_w, img_h)

    f_box = s13.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Sơ đồ kiến trúc kiểm thử đối kháng do sinh viên tự thiết kế minh họa các luận điểm của Halpin (2025)."
    p_f.font.size = Pt(10)
    p_f.font.italic = True
    p_f.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 14: KẾT LUẬN & ĐỊNH HƯỚNG BÁO CÁO GIỮA KỲ
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_bg(s14)
    add_header(s14, "TỔNG KẾT BÁO CÁO GIỮA KỲ & ĐỊNH HƯỚNG NGHIÊN CỨU", "SECTION 7 — KẾT LUẬN & THÔNG ĐIỆP HỌC THUẬT",
               "Tổng kết các đóng góp lý thuyết của Halpin và lộ trình kỹ thuật cho giai đoạn cuối kỳ")

    # Left: Đóng góp bài báo Halpin
    add_card(s14, Inches(0.8), Inches(1.85), Inches(5.7), Inches(4.6), COLOR_BG_PANEL, COLOR_BORDER_BLUE)
    t = s14.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.3), Inches(4.3))
    tf = t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "1. Thông Điệp Học Thuật Cốt Lõi (Harry Halpin 2025)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_BRAND_BLUE

    conclusions = [
        ("Vượt qua sự trừu tượng hóa:", "Mật mã học không chỉ là toán học thuần túy trên giấy tờ. Các hệ thống an toàn trong thế giới thực phải nhận thức rõ khoảng cách giữa mô hình lý tưởng và thực tế vật lý (kênh kề, siêu dữ liệu, chuỗi cung ứng, áp lực thể chế)."),
        ("Bản chất của đối thủ:", "Đối thủ không phải là 'nhiễu ngẫu nhiên' mà là tác nhân có chủ đích, liên tục tìm kiếm các kênh rò rỉ bị mô hình toán học bỏ qua."),
        ("Chương trình nghiên cứu mở:", "Halpin đặt nền móng triết học và vạch ra 4 câu hỏi nghiên cứu liên ngành nhằm mở rộng mật mã học thành một khoa học an ninh hoàn chỉnh.")
    ]
    for h, b in conclusions:
        p_it = tf.add_paragraph()
        p_it.space_before = Pt(8)
        r1 = p_it.add_run()
        r1.text = f"• {h} "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_TEXT_PRIMARY
        r2 = p_it.add_run()
        r2.text = b
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = COLOR_TEXT_SECONDARY

    # Right: Kết quả giữa kỳ & Định hướng cuối kỳ
    add_card(s14, Inches(6.8), Inches(1.85), Inches(5.7), Inches(4.6), COLOR_EMERALD_BG, COLOR_EMERALD_BORDER)
    t = s14.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.3))
    tf = t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "2. Hoàn Thành Giữa Kỳ & Lộ Trình Cuối Kỳ (Nhóm DT005)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT_EMERALD

    next_steps = [
        ("Nghiệm thu tiến độ giữa kỳ:", "Hoàn thành thẩm định học thuật sâu sắc toàn văn bài báo gốc (17 trang PDF), phân tích đầy đủ 7 mục nội dung và xây dựng bộ diễn tập minh họa thực nghiệm sống động phục vụ báo cáo."),
        ("Minh bạch phương pháp luận:", "Không ngụy tạo số liệu tác giả; phân định rạch ròi giữa công trình triết học của Halpin và các thử nghiệm độc lập của sinh viên."),
        ("Lộ trình nghiên cứu cuối kỳ:", "Đề xuất phát triển chuyên sâu giải pháp tối ưu hóa cấu trúc cây Merkle phân cấp (HD-Merkle Dirty-Path Caching, nén CompactProof O(1)) làm đóng góp kỹ thuật trọng tâm cho đồ án cuối kỳ.")
    ]
    for h, b in next_steps:
        p_it = tf.add_paragraph()
        p_it.space_before = Pt(8)
        r1 = p_it.add_run()
        r1.text = f"• {h} "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_TEXT_PRIMARY
        r2 = p_it.add_run()
        r2.text = b
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = COLOR_TEXT_SECONDARY

    # Footer nhỏ gọn ghi chú tài liệu tham khảo chính
    f_box = s14.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.55))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "Tài liệu tham khảo chính: [1] H. Halpin, Oxford J. Cybersec., 2025 (DOI: 10.1093/cybsec/tyaf006); [2] Shannon (1949); [3] Goldwasser-Micali (1984); [4] Dolev-Yao (1983); [5] Kocher (1996); [6] RFC 6962 (2013)."
    p_f.font.size = Pt(9.5)
    p_f.font.italic = True
    p_f.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 15: SLIDE CẢM ƠN & Q&A (ĐỒNG BỘ THIẾT KẾ 100% VỚI SLIDE BÌA)
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    add_bg(s15)

    # Dải màu xanh dương dọc bên lề trái đồng bộ nhận diện
    left_bar15 = s15.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.8), Inches(0.18), Inches(5.9))
    left_bar15.fill.solid()
    left_bar15.fill.fore_color.rgb = COLOR_ACCENT_BLUE
    left_bar15.line.fill.background()

    # Khối tiêu đề Trường & Khoa ở phần trên
    t_inst15 = s15.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11.0), Inches(0.9))
    tf_inst15 = t_inst15.text_frame
    tf_inst15.word_wrap = True
    p_uni15 = tf_inst15.paragraphs[0]
    p_uni15.alignment = PP_ALIGN.CENTER
    p_uni15.text = "Trường Đại học Tài nguyên và Môi trường TP. Hồ Chí Minh"
    p_uni15.font.size = Pt(15)
    p_uni15.font.bold = True
    p_uni15.font.color.rgb = COLOR_TEXT_PRIMARY

    p_fac15 = tf_inst15.add_paragraph()
    p_fac15.alignment = PP_ALIGN.CENTER
    p_fac15.space_before = Pt(4)
    p_fac15.text = "Khoa Công nghệ thông tin"
    p_fac15.font.size = Pt(14)
    p_fac15.font.bold = True
    p_fac15.font.color.rgb = COLOR_TEXT_SECONDARY

    # Khối cảm ơn & Q&A ở giữa
    t_thank = s15.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(11.0), Inches(2.0))
    tf_thank = t_thank.text_frame
    tf_thank.word_wrap = True

    p_thk = tf_thank.paragraphs[0]
    p_thk.alignment = PP_ALIGN.CENTER
    p_thk.text = "CẢM ƠN THẦY VÀ CÁC BẠN ĐÃ LẮNG NGHE!"
    p_thk.font.size = Pt(32)
    p_thk.font.bold = True
    p_thk.font.color.rgb = COLOR_ACCENT_BLUE

    p_qa = tf_thank.add_paragraph()
    p_qa.alignment = PP_ALIGN.CENTER
    p_qa.space_before = Pt(14)
    p_qa.text = "Q & A — Xin đón nhận các câu hỏi và ý kiến đóng góp của Thầy và Hội đồng"
    p_qa.font.size = Pt(16)
    p_qa.font.italic = True
    p_qa.font.color.rgb = COLOR_TEXT_SECONDARY

    # Khối thông tin đề tài & sinh viên ở góc dưới
    t_info15 = s15.shapes.add_textbox(Inches(1.4), Inches(4.9), Inches(10.0), Inches(1.8))
    tf_info15 = t_info15.text_frame
    tf_info15.word_wrap = True

    p_t15 = tf_info15.paragraphs[0]
    p_t15.text = "Đề tài DT005: Hệ thống kiểm tra toàn vẹn tệp bằng SHA-256, HMAC và cây Merkle phân cấp động"
    p_t15.font.size = Pt(13)
    p_t15.font.bold = True
    p_t15.font.color.rgb = COLOR_TEXT_TITLE

    p_c15 = tf_info15.add_paragraph()
    p_c15.space_before = Pt(4)
    p_c15.text = "Học phần: An toàn và bảo mật hệ thống thông tin"
    p_c15.font.size = Pt(12.5)
    p_c15.font.color.rgb = COLOR_TEXT_PRIMARY

    p_lec15 = tf_info15.add_paragraph()
    p_lec15.space_before = Pt(4)
    p_lec15.text = "Giảng viên hướng dẫn: ThS. Phạm Trọng Huynh"
    p_lec15.font.size = Pt(12.5)
    p_lec15.font.color.rgb = COLOR_TEXT_PRIMARY

    p_stu15 = tf_info15.add_paragraph()
    p_stu15.space_before = Pt(4)
    r15_1 = p_stu15.add_run()
    r15_1.text = "Sinh viên thực hiện: "
    r15_1.font.size = Pt(12.5)
    r15_1.font.color.rgb = COLOR_TEXT_PRIMARY
    r15_2 = p_stu15.add_run()
    r15_2.text = "Võ Anh Kiệt"
    r15_2.font.bold = True
    r15_2.font.size = Pt(12.5)
    r15_2.font.color.rgb = COLOR_BRAND_BLUE
    r15_3 = p_stu15.add_run()
    r15_3.text = " — MSSV: 1150080143 — Lớp: 11_CNPM2"
    r15_3.font.size = Pt(12.5)
    r15_3.font.color.rgb = COLOR_TEXT_PRIMARY

    # Lưu file
    out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "docs")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Bao_Cao_Giua_Ky_DT05_Final.pptx")
    prs.save(out_path)
    
    root_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Bao_Cao_Giua_Ky_DT05_Final.pptx")
    prs.save(root_path)

    print(f"[THÀNH CÔNG] Đã tạo bài thuyết trình giữa kỳ hoàn chỉnh tại:\n  1. {out_path}\n  2. {root_path}")
    print(f"Tổng số slide: {len(prs.slides)} slide chuẩn 16:9 widescreen.")


if __name__ == "__main__":
    create_deck()
